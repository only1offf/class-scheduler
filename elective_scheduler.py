#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
선택과목 이동반 편성 프로그램
Usage:
  python elective_scheduler.py [CSV파일] [--homeroom 본반CSV] [--names 이름CSV] [--seed N] [--out 출력파일.xlsx]
"""
import os
import sys
import random

# Windows 콘솔 UTF-8 출력
if sys.stdout.encoding and sys.stdout.encoding.lower() not in ('utf-8', 'utf8'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import argparse
from collections import defaultdict
from itertools import permutations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

try:
    import pulp
    HAS_PULP = True
except ImportError:
    HAS_PULP = False

# ──────────────────────────────────────────────────────────────────────────────
# 상수 정의
# ──────────────────────────────────────────────────────────────────────────────

SUBJECTS = [
    '기하', '동아시아역사기행', '법과사회', '세계시민과지리', '정치',
    '현대사회와윤리', '물질과에너지', '세포와물질대사', '역학과에너지', '지구시스템과학',
]

N_SECTIONS = {
    '기하': 3, '동아시아역사기행': 2, '법과사회': 3, '세계시민과지리': 3,
    '정치': 2, '현대사회와윤리': 5, '물질과에너지': 4,
    '세포와물질대사': 3, '역학과에너지': 3, '지구시스템과학': 2,
}

# 타임당 최대 동시 분반 수 (= 교사 수)
MAX_PER_TIME = {
    '기하': 2, '현대사회와윤리': 2, '물질과에너지': 2,
    '동아시아역사기행': 1, '법과사회': 1, '세계시민과지리': 1,
    '정치': 1, '세포와물질대사': 1, '역학과에너지': 1, '지구시스템과학': 1,
}

MAX_PER_SECTION = 41
TIMES = ['A', 'B', 'C']
SECTIONS_PER_TIME = 10


def extract_homeroom(sid):
    """신학번 5자리에서 본반 추출: 20103 → '1반', 20815 → '8반'"""
    s = str(sid).strip()
    if len(s) == 5:
        return str(int(s[1:3])) + '반'
    return ''

ABBR = {
    '기하': '기하', '동아시아역사기행': '동아역사', '법과사회': '법과사회',
    '세계시민과지리': '세계지리', '정치': '정치', '현대사회와윤리': '현대사회',
    '물질과에너지': '물질에너지', '세포와물질대사': '세포대사',
    '역학과에너지': '역학에너지', '지구시스템과학': '지구시스템',
}

# ──────────────────────────────────────────────────────────────────────────────
# 데이터 로드
# ──────────────────────────────────────────────────────────────────────────────

def load_data(path):
    df = pd.read_csv(path, encoding='utf-8-sig', dtype=str)
    df.columns = df.columns.str.strip()
    for col in df.columns:
        df[col] = df[col].str.strip()
    for s in SUBJECTS:
        if s in df.columns:
            df[s] = pd.to_numeric(df[s], errors='coerce').fillna(0).astype(int)
    return df


def load_extra_csv(path, key_col, val_col):
    """key_col → val_col 매핑 딕셔너리 반환"""
    if not path or not os.path.exists(path):
        return {}
    df = pd.read_csv(path, encoding='utf-8-sig', dtype=str)
    df.columns = df.columns.str.strip()
    df[key_col] = df[key_col].str.strip()
    df[val_col] = df[val_col].str.strip()
    return dict(zip(df[key_col], df[val_col]))

# ──────────────────────────────────────────────────────────────────────────────
# 데이터 분석 출력
# ──────────────────────────────────────────────────────────────────────────────

def print_analysis(df):
    sep = "=" * 65
    print(sep)
    print("■ CSV 구조 분석")
    print(f"  총 학생 수: {len(df)}명  |  컬럼 수: {len(df.columns)}")
    print(f"  컬럼: {', '.join(df.columns.tolist())}")
    print()

    print("■ 과목별 신청 인원 및 분반 정보")
    print(f"  {'과목':<20} {'신청':>6} {'분반':>5} {'교사(max/타임)':>14} {'평균':>7} {'41명OK':>8}")
    print("  " + "─" * 62)
    for s in SUBJECTS:
        cnt = int(df[s].sum())
        n   = N_SECTIONS[s]
        avg = cnt / n
        ok  = "✓" if avg <= MAX_PER_SECTION else "★초과★"
        print(f"  {s:<20} {cnt:>6} {n:>5} {MAX_PER_TIME[s]:>14} {avg:>7.1f} {ok:>8}")

    total_enroll = sum(int(df[s].sum()) for s in SUBJECTS)
    print(f"  {'합계':<20} {total_enroll:>6} {sum(N_SECTIONS.values()):>5}")
    print()

    bad = df[df[SUBJECTS].sum(axis=1) != 3]
    if len(bad):
        print(f"  ★ 3과목 미준수 학생 {len(bad)}명: {bad['신학번'].tolist()}")
    else:
        print("  ✓ 모든 학생이 정확히 3과목 선택")

    from collections import Counter
    combos = Counter(
        tuple(s for s in SUBJECTS if row[s] == 1)
        for _, row in df.iterrows()
    )
    print(f"  고유 과목 조합: {len(combos)}가지")
    print()
    print("  [상위 5개 조합]")
    for combo, cnt in combos.most_common(5):
        print(f"    {'+'.join(ABBR[s] for s in combo)}: {cnt}명")
    print(sep)
    print()

# ──────────────────────────────────────────────────────────────────────────────
# Phase 1: 분반 → 타임 배정 (ILP 또는 기본 배분)
# ──────────────────────────────────────────────────────────────────────────────

def solve_with_ilp():
    """PuLP ILP로 각 과목의 타임별 분반 수 결정"""
    prob = pulp.LpProblem("section_time", pulp.LpMinimize)

    z = {
        s: {t: pulp.LpVariable(f"z_{i}_{t}", 0, N_SECTIONS[s], cat='Integer')
            for t in TIMES}
        for i, s in enumerate(SUBJECTS)
    }

    prob += 0  # 실현 가능한 해만 찾으면 됨

    for s in SUBJECTS:
        prob += pulp.lpSum(z[s][t] for t in TIMES) == N_SECTIONS[s]
        for t in TIMES:
            prob += z[s][t] <= MAX_PER_TIME[s]

    for t in TIMES:
        prob += pulp.lpSum(z[s][t] for s in SUBJECTS) == SECTIONS_PER_TIME

    prob.solve(pulp.PULP_CBC_CMD(msg=0))

    if pulp.LpStatus[prob.status] != 'Optimal':
        return None

    return {s: {t: int(round(pulp.value(z[s][t]))) for t in TIMES} for s in SUBJECTS}


def fallback_assignment():
    """ILP 없이 사전 검증된 배분 사용 (전원 충족 가능 확인)
    타임별 합 = 10, 교사 제약 충족, 57개 전 조합 Hall 조건 통과.
    """
    return {
        '기하':             {'A': 1, 'B': 1, 'C': 1},
        '동아시아역사기행':  {'A': 1, 'B': 1, 'C': 0},
        '법과사회':          {'A': 1, 'B': 1, 'C': 1},
        '세계시민과지리':    {'A': 1, 'B': 1, 'C': 1},
        '정치':              {'A': 1, 'B': 0, 'C': 1},
        '현대사회와윤리':    {'A': 1, 'B': 2, 'C': 2},
        '물질과에너지':      {'A': 2, 'B': 1, 'C': 1},
        '세포와물질대사':    {'A': 1, 'B': 1, 'C': 1},
        '역학과에너지':      {'A': 1, 'B': 1, 'C': 1},
        '지구시스템과학':    {'A': 0, 'B': 1, 'C': 1},
    }


def solve_section_assignment():
    if HAS_PULP:
        result = solve_with_ilp()
        if result:
            return result, 'ILP(PuLP)'
    return fallback_assignment(), '기본배분(사전검증)'


def build_section_list(assignment):
    """과목별로 분반번호(1-based, 전체통합) → 타임 매핑 생성
    returns:
      sections: {과목: {타임: [분반번호, ...]}}
      sec_to_time: {(과목, 분반번호): 타임}
    """
    sections = {}
    sec_to_time = {}
    for s in SUBJECTS:
        sections[s] = {}
        num = 1
        for t in TIMES:
            cnt = assignment[s][t]
            secs = list(range(num, num + cnt))
            sections[s][t] = secs
            for k in secs:
                sec_to_time[(s, k)] = t
            num += cnt
    return sections, sec_to_time

# ──────────────────────────────────────────────────────────────────────────────
# Phase 2: 학생 → 분반 배정 (탐욕 이분 매칭)
# ──────────────────────────────────────────────────────────────────────────────

def assign_students(df, assignment, sections, names=None, homeroom=None):
    names    = names or {}
    homeroom = homeroom or {}

    # (과목, 분반번호) → 현재 인원
    sec_load = {(s, k): 0 for s in SUBJECTS for t in TIMES for k in sections[s][t]}
    sec_roster = defaultdict(list)   # (과목, 분반번호) → [신학번, ...]
    records = []
    failed  = []

    rows = df.to_dict('records')
    random.shuffle(rows)

    for row in rows:
        sid    = str(row['신학번'])
        chosen = [s for s in SUBJECTS if int(row.get(s, 0)) == 1]

        if len(chosen) != 3:
            failed.append(sid)
            continue

        # 6가지 타임 순열 중 최적 배정 선택
        # 우선순위 1: 모든 분반에 여석(< 41명)이 있는 순열 중 최소 부하
        # 우선순위 2: 여석 없으면 부하 최소 순열 (초과 최소화 fallback)
        best_map      = None   # 여석 있는 최적
        best_cost     = float('inf')
        overflow_map  = None   # fallback (여석 없어도 부하 최소)
        overflow_cost = float('inf')

        for perm in permutations(TIMES):
            tm = dict(zip(chosen, perm))
            if not all(assignment[s][t] > 0 for s, t in tm.items()):
                continue
            cost = sum(min(sec_load[(s, k)] for k in sections[s][tm[s]]) for s in chosen)
            # 각 과목의 해당 타임에 여석 있는 분반이 존재하는지 확인
            has_room = all(
                any(sec_load[(s, k)] < MAX_PER_SECTION for k in sections[s][tm[s]])
                for s in chosen
            )
            if has_room and cost < best_cost:
                best_cost, best_map = cost, tm
            if cost < overflow_cost:
                overflow_cost, overflow_map = cost, tm

        final_map = best_map if best_map is not None else overflow_map
        if final_map is None:
            failed.append(sid)
            continue

        rec = {
            '신학번': sid,
            '이름':   names.get(sid, ''),
            '본반':   homeroom.get(sid, ''),
        }
        # 타임 컬럼 초기화
        for t in TIMES:
            rec[f'{t}타임_과목'] = ''
            rec[f'{t}타임_분반'] = ''

        for s in chosen:
            t = final_map[s]
            # 여석 있는 분반 우선, 없으면 가장 덜 찬 분반
            with_room = [k for k in sections[s][t] if sec_load[(s, k)] < MAX_PER_SECTION]
            pool = with_room if with_room else sections[s][t]
            k = min(pool, key=lambda x: sec_load[(s, x)])
            sec_load[(s, k)]  += 1
            sec_roster[(s, k)].append(sid)
            rec[f'{t}타임_과목'] = s
            rec[f'{t}타임_분반'] = f"{ABBR[s]}-{k}반"

        records.append(rec)

    rec_df = pd.DataFrame(records) if records else pd.DataFrame()

    # 원본 순서 복원
    base = df[['신학번']].copy()
    base['신학번'] = base['신학번'].astype(str)
    result_df = base.merge(rec_df, on='신학번', how='left') if not rec_df.empty else base

    return result_df, sec_roster, sec_load, failed

# ──────────────────────────────────────────────────────────────────────────────
# 검증 출력
# ──────────────────────────────────────────────────────────────────────────────

def print_validation(assignment, sections, sec_load, failed, n_total, method):
    sep = "=" * 65
    print(sep)
    print(f"■ 배정 결과 검증  [방법: {method}]")
    print()

    print("  [타임별 분반 배정]")
    print(f"  {'과목':<20} {'A':>3} {'B':>3} {'C':>3}   분반 목록")
    print("  " + "─" * 60)
    for s in SUBJECTS:
        a, b, c = assignment[s]['A'], assignment[s]['B'], assignment[s]['C']
        tags = [f"{t}{k}" for t in TIMES for k in sections[s][t]]
        print(f"  {s:<20} {a:>3} {b:>3} {c:>3}   {' '.join(tags)}")
    tots = [sum(assignment[s][t] for s in SUBJECTS) for t in TIMES]
    print(f"  {'합계':<20} {tots[0]:>3} {tots[1]:>3} {tots[2]:>3}")
    print()

    print("  [분반별 배정 인원]  (★=41명 초과)")
    over_count = 0
    for s in SUBJECTS:
        parts = []
        for t in TIMES:
            for k in sections[s][t]:
                cnt  = sec_load.get((s, k), 0)
                flag = "★" if cnt > MAX_PER_SECTION else ""
                parts.append(f"{t}{k}반:{cnt}{flag}")
                if cnt > MAX_PER_SECTION:
                    over_count += 1
        print(f"  {s:<20} {' | '.join(parts)}")

    print()
    ok = n_total - len(failed)
    print(f"  학생 충족률: {ok}/{n_total}명 ({100 * ok / n_total:.1f}%)")
    if failed:
        print(f"  미배정 학생({len(failed)}명): {failed[:10]}{'...' if len(failed) > 10 else ''}")
    if over_count:
        print(f"  41명 초과 분반: {over_count}개")
    else:
        print("  ✓ 모든 분반 41명 이하")
    print(sep)
    print()

# ──────────────────────────────────────────────────────────────────────────────
# Excel 출력 헬퍼
# ──────────────────────────────────────────────────────────────────────────────

_HDR_FILL  = PatternFill("solid", fgColor="2E4699")
_HDR_FONT  = Font(bold=True, color="FFFFFF", size=10)
_HDR_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
_SUB_FILL  = PatternFill("solid", fgColor="5B9BD5")
_SUB_FONT  = Font(bold=True, color="FFFFFF", size=10)
_TIT_FONT  = Font(bold=True, size=12)


def hdr(ws, row, col, value):
    c = ws.cell(row, col, value)
    c.font, c.fill, c.alignment = _HDR_FONT, _HDR_FILL, _HDR_ALIGN
    return c


def subhdr(ws, row, col, value):
    c = ws.cell(row, col, value)
    c.font, c.fill, c.alignment = _SUB_FONT, _SUB_FILL, _HDR_ALIGN
    return c


def auto_width(ws, min_w=8, max_w=40, extra=2):
    for col in ws.columns:
        best = max((len(str(c.value or '')) for c in col), default=0)
        ws.column_dimensions[
            get_column_letter(col[0].column)
        ].width = min(max(best + extra, min_w), max_w)

# ──────────────────────────────────────────────────────────────────────────────
# Excel 시트 1: 학생시간표
# ──────────────────────────────────────────────────────────────────────────────

def write_student_timetable(wb, result_df):
    ws = wb.create_sheet("학생시간표")
    cols = [
        '신학번', '이름', '본반',
        'A타임_과목', 'A타임_분반',
        'B타임_과목', 'B타임_분반',
        'C타임_과목', 'C타임_분반',
    ]
    for ci, h in enumerate(cols, 1):
        hdr(ws, 1, ci, h)
    ws.row_dimensions[1].height = 20

    for _, row in result_df.iterrows():
        ws.append([str(row.get(c, '') or '') for c in cols])

    ws.freeze_panes = 'A2'
    auto_width(ws)

# ──────────────────────────────────────────────────────────────────────────────
# Excel 시트 2: 분반별 명단
# ──────────────────────────────────────────────────────────────────────────────

def write_section_rosters(wb, sec_roster, sections, names, homeroom):
    for s in SUBJECTS:
        for t in TIMES:
            for k in sections[s][t]:
                ws = wb.create_sheet(f"{ABBR[s]}-{k}반")

                # 제목 행
                title = f"[{s}]  {k}분반  ({t}타임)"
                ws.cell(1, 1, title).font = _TIT_FONT
                ws.merge_cells('A1:D1')
                ws.row_dimensions[1].height = 22

                # 헤더
                for ci, h in enumerate(['신학번', '이름', '본반', '비고'], 1):
                    hdr(ws, 2, ci, h)
                ws.row_dimensions[2].height = 18

                # 데이터 (신학번 기준 정렬)
                for sid in sorted(sec_roster.get((s, k), [])):
                    ws.append([sid, names.get(sid, ''), homeroom.get(sid, ''), ''])

                auto_width(ws)

# ──────────────────────────────────────────────────────────────────────────────
# Excel 시트 3: 개요 (타임별 분반 현황)
# ──────────────────────────────────────────────────────────────────────────────

def write_summary(wb, assignment, sections, sec_load, n_total, failed):
    ws = wb.create_sheet("개요", 0)  # 첫 번째 시트

    ws.cell(1, 1, "이동반 편성 개요").font = Font(bold=True, size=14)
    ws.merge_cells('A1:G1')
    ws.row_dimensions[1].height = 25

    ws.cell(2, 1, f"총 학생: {n_total}명  |  충족: {n_total - len(failed)}명  |  미배정: {len(failed)}명")
    ws.merge_cells('A2:G2')

    row = 4
    for ci, h in enumerate(['타임', '과목', '분반명', '분반번호', '배정인원', '최대인원', '상태'], 1):
        hdr(ws, row, ci, h)
    ws.row_dimensions[row].height = 18
    row += 1

    for t in TIMES:
        subhdr(ws, row, 1, f"{t}타임")
        ws.merge_cells(f'A{row}:G{row}')
        ws.row_dimensions[row].height = 16
        row += 1

        for s in SUBJECTS:
            for k in sections[s][t]:
                cnt   = sec_load.get((s, k), 0)
                state = "★초과★" if cnt > MAX_PER_SECTION else "정상"
                ws.cell(row, 1, t)
                ws.cell(row, 2, s)
                ws.cell(row, 3, f"{ABBR[s]}-{k}반")
                ws.cell(row, 4, k)
                ws.cell(row, 5, cnt)
                ws.cell(row, 6, MAX_PER_SECTION)
                ws.cell(row, 7, state)
                if state != "정상":
                    ws.cell(row, 7).font = Font(bold=True, color="FF0000")
                row += 1

        row += 1  # 타임 간 빈 줄

    auto_width(ws)

# ──────────────────────────────────────────────────────────────────────────────
# Excel 시트 4: 교실 배치 추천
# ──────────────────────────────────────────────────────────────────────────────

def write_classroom_recommendation(wb, assignment, sections, sec_roster, sec_load, homeroom):
    """타임별로 각 분반의 본반 구성을 보여주는 교실배치 추천 시트"""
    ws = wb.create_sheet("교실배치추천")

    ws.cell(1, 1, "교실 배치 추천 — 본반 우선순위 기준").font = _TIT_FONT
    ws.merge_cells('A1:F1')
    ws.row_dimensions[1].height = 22
    ws.cell(2, 1, "※ '본반 구성(우선순위)'은 해당 분반 학생의 본반 분포를 인원 많은 순으로 표시. 교실 배정 시 참고.")
    ws.cell(2, 1).font = Font(italic=True, size=9, color="595959")
    ws.merge_cells('A2:F2')

    row = 4
    for t in TIMES:
        # 타임 구분 헤더
        subhdr(ws, row, 1, f"▶ {t}타임")
        ws.merge_cells(f'A{row}:F{row}')
        ws.row_dimensions[row].height = 18
        row += 1

        col_hdrs = ['과목', '분반명', '배정인원', '본반 구성 (우선순위)', '권장 교실', '비고']
        for ci, h in enumerate(col_hdrs, 1):
            hdr(ws, row, ci, h)
        ws.row_dimensions[row].height = 18
        row += 1

        for s in SUBJECTS:
            for k in sections[s][t]:
                cnt = sec_load.get((s, k), 0)

                # 본반별 학생 수 집계 → 많은 순 정렬
                hr_cnt = defaultdict(int)
                for sid in sec_roster.get((s, k), []):
                    hr = homeroom.get(sid, '?')
                    hr_cnt[hr] += 1
                top = sorted(hr_cnt.items(),
                             key=lambda x: (-x[1], x[0]))  # 인원 내림, 본반명 오름

                # "1반 12명 > 2반 10명 > 3반 8명 > ..." 형식
                priority_str = '  >  '.join(
                    f"{hr} {c}명" for hr, c in top
                )

                ws.cell(row, 1, s)
                ws.cell(row, 2, f"{ABBR[s]}-{k}반")
                ws.cell(row, 3, cnt)
                ws.cell(row, 4, priority_str)
                ws.cell(row, 5, '')   # 권장 교실 (관리자 직접 입력)
                ws.cell(row, 6, '')
                ws.cell(row, 4).alignment = Alignment(horizontal='left')
                row += 1

        row += 1  # 타임 간 빈 줄

    # 컬럼 너비 수동 조정 (본반 구성 컬럼이 매우 길어질 수 있음)
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 60
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 14

# ──────────────────────────────────────────────────────────────────────────────
# 메인
# ──────────────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description='선택과목 이동반 편성 프로그램')
    parser.add_argument(
        'csv', nargs='?',
        default=r'C:\Users\User\class-scheduler\학생선택과목.csv',
        help='학생 선택과목 CSV 파일',
    )
    parser.add_argument('--homeroom', help='본반 CSV (컬럼: 신학번,본반)')
    parser.add_argument('--names',    help='이름 CSV (컬럼: 신학번,이름)')
    parser.add_argument('--seed',     type=int, default=42, help='무작위 시드 (기본 42)')
    parser.add_argument('--out',      default=None, help='출력 Excel 파일명')
    args = parser.parse_args()

    random.seed(args.seed)

    bar = "─" * 65
    print(f"\n{bar}")
    print("  선택과목 이동반 편성 프로그램")
    print(f"{bar}\n")

    print(f"▶ 데이터 로드: {args.csv}")
    df    = load_data(args.csv)
    names = load_extra_csv(args.names, '신학번', '이름') if args.names else {}

    # 신학번에서 본반 자동 추출 (5자리: [학년][반2자리][번호2자리])
    homeroom = {str(row['신학번']): extract_homeroom(str(row['신학번']))
                for _, row in df.iterrows()}
    # --homeroom CSV 제공 시 덮어씀
    if args.homeroom:
        homeroom.update(load_extra_csv(args.homeroom, '신학번', '본반'))

    unique_hr = sorted(set(homeroom.values()), key=lambda x: int(x.replace('반', '')))
    print(f"  본반 자동 추출: {len(unique_hr)}개 반 ({', '.join(unique_hr)})")
    if names:
        print(f"  이름 데이터: {len(names)}명")
    print()

    print_analysis(df)

    print("▶ 분반 타임 배정 중...")
    assignment, method = solve_section_assignment()
    sections, sec_to_time = build_section_list(assignment)
    print(f"  완료 [{method}]\n")

    print("▶ 학생 분반 배정 중...")
    result_df, sec_roster, sec_load, failed = assign_students(
        df, assignment, sections, names, homeroom,
    )
    print_validation(assignment, sections, sec_load, failed, len(df), method)

    out_dir  = os.path.dirname(os.path.abspath(args.csv))
    out_xlsx = args.out or os.path.join(out_dir, '이동반편성결과.xlsx')
    out_csv  = os.path.join(out_dir, '학생시간표.csv')

    print("▶ Excel 생성 중...")
    wb = Workbook()
    wb.remove(wb.active)

    write_summary(wb, assignment, sections, sec_load, len(df), failed)
    write_student_timetable(wb, result_df)
    write_section_rosters(wb, sec_roster, sections, names, homeroom)
    write_classroom_recommendation(wb, assignment, sections, sec_roster, sec_load, homeroom)

    wb.save(out_xlsx)
    print(f"  Excel 저장: {out_xlsx}")

    result_df.to_csv(out_csv, index=False, encoding='utf-8-sig')
    print(f"  CSV 저장:   {out_csv}")
    print(f"\n✓ 완료!\n{bar}\n")


if __name__ == '__main__':
    main()
